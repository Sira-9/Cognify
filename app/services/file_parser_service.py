import os
from pathlib import Path
from typing import Dict, Any, Tuple
from app.config import MAX_FILE_SIZE_MB, ALLOWED_EXTENSIONS

class FileValidationError(Exception):
    pass

def validate_uploaded_file(filename: str, file_size: int) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise FileValidationError(
            f"Unsupported file type '{ext}'. Please upload a PDF, DOC/DOCX, PPT/PPTX, or TXT file."
        )
    
    max_bytes = MAX_FILE_SIZE_MB * 1024 * 1024
    if file_size > max_bytes:
        raise FileValidationError(
            f"File is too large ({file_size / (1024*1024):.1f} MB). Maximum allowed size is {MAX_FILE_SIZE_MB} MB."
        )
    
    if file_size <= 0:
        raise FileValidationError("The uploaded file is empty (0 bytes).")
    
    return ext

def parse_txt_file(file_path: Path) -> str:
    for encoding in ['utf-8', 'latin-1', 'cp1252']:
        try:
            with open(file_path, 'r', encoding=encoding) as f:
                content = f.read()
                if content.strip():
                    return content
        except UnicodeDecodeError:
            continue
    raise FileValidationError("Could not decode TXT file. Ensure it contains valid text.")

def parse_pdf_file(file_path: Path) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(file_path))
        if len(reader.pages) == 0:
            raise FileValidationError("PDF has no pages.")
        
        extracted_text = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                extracted_text.append(f"--- Page {i+1} ---\n" + text.strip())
        
        full_text = "\n\n".join(extracted_text)
        if len(full_text.strip()) < 20:
            raise FileValidationError("No readable text found in the PDF. It may be scanned images or corrupted.")
        return full_text
    except FileValidationError:
        raise
    except Exception as e:
        raise FileValidationError(f"Corrupted or invalid PDF file: {str(e)}")

def parse_docx_file(file_path: Path) -> str:
    try:
        import docx
        doc = docx.Document(str(file_path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        
        full_text = "\n\n".join(paragraphs)
        if len(full_text.strip()) < 20:
            raise FileValidationError("DOCX file contains no readable text paragraphs.")
        return full_text
    except FileValidationError:
        raise
    except Exception as e:
        raise FileValidationError(f"Could not parse Word document: {str(e)}")

def parse_pptx_file(file_path: Path) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(str(file_path))
        slide_texts = []
        for i, slide in enumerate(prs.slides):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.text.strip():
                            texts.append(paragraph.text.strip())
            if texts:
                slide_texts.append(f"--- Slide {i+1} ---\n" + "\n".join(texts))
        
        full_text = "\n\n".join(slide_texts)
        if len(full_text.strip()) < 20:
            raise FileValidationError("Presentation contains no readable slide text.")
        return full_text
    except FileValidationError:
        raise
    except Exception as e:
        raise FileValidationError(f"Could not parse PowerPoint presentation: {str(e)}")

def extract_content_from_file(file_path: Path, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    
    if ext == ".txt":
        text = parse_txt_file(file_path)
    elif ext == ".pdf":
        text = parse_pdf_file(file_path)
    elif ext in [".docx", ".doc"]:
        text = parse_docx_file(file_path)
    elif ext in [".pptx", ".ppt"]:
        text = parse_pptx_file(file_path)
    else:
        raise FileValidationError(f"Unsupported file format: {ext}")
    
    # Final check on meaningful content
    cleaned = text.strip()
    if len(cleaned) < 30:
        raise FileValidationError("The document did not contain sufficient textual content to build a study course.")
    
    return cleaned
